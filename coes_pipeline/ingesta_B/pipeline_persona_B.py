#!/usr/bin/env python3
"""
Pipeline Persona B - Eventos, Confiabilidad y Reprograma (COES / SEIN)

Genera 4 tablas en formato largo + notas de calidad:
  B1_mantenimientos.csv            (deduplicado)
  B2_potencia_indisponible.csv
  B3_reprograma_diario.csv         (incluye motivo_categoria y horas_cubiertas)
  B4_pr06_rdo.csv                  (label RDO + id_reprograma, UNA fila por bloque)
  B4_pr06_rdo_por_archivo.csv      (crudo: una fila por bloque y por archivo/letra)
  B4_consistencia_rdo_vs_forzado.csv
  B4_diagnostico_abg.csv           (que bloques de alfa/beta/gamma cambian entre letras)
  B_dataset_maestro.csv            (UNA fila por bloque: label RDO + features de B1/B2 + columnas de impacto de B3)
  NOTAS_B.txt

Uso:
  python pipeline_persona_B.py --mant RptMantenimiento.xlsx \
      --pot-dir datos/potencia --reprog-dir datos/reprograma \
      --pr06-dir datos/pr06 --out salida/

Convenciones acordadas con Persona A:
  * Ventana: 2021-12-01 a 2024-11-30
  * Timestamp unico: fecha_hora 'AAAA-MM-DD HH:MM'
  * Bloques de 30 min con convencion "hora fin de bloque" (00:30 ... 00:00 del dia siguiente)

SUPUESTOS SOBRE LOS ARCHIVOS CRUDOS (verificar con una muestra real, ver NOTAS_B.txt):
  * Los encabezados se localizan buscando palabras clave, no posiciones fijas.
  * La fecha de B2 se toma del nombre de archivo (AAAAMMDD o AAAA-MM-DD).
  * La fecha de B3 se toma del nombre de la hoja; la letra, del nombre del archivo.
  * El PR06 es un archivo por dia (indicesAAAAMMDD.xlsx) con columnas Hora, RDO, alfa, beta, gamma, CHM.
"""
import argparse
import re
import unicodedata
from datetime import datetime, time, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

VENTANA_INI = pd.Timestamp("2021-12-01 00:00")
VENTANA_FIN = pd.Timestamp("2024-11-30 23:59")
FMT = "%Y-%m-%d %H:%M"
NOTAS = []


def nota(msg):
    print(msg)
    NOTAS.append(msg)


def norm(s):
    """MAYUSCULAS sin tildes ni espacios sobrantes."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def en_ventana(ts):
    return (ts >= VENTANA_INI) & (ts <= VENTANA_FIN)


def parse_fecha_archivo(nombre):
    """AAAAMMDD o AAAA-MM-DD dentro de un nombre de archivo."""
    m = re.search(r"(20\d{2})[-_]?(\d{2})[-_]?(\d{2})", nombre)
    return pd.Timestamp(f"{m.group(1)}-{m.group(2)}-{m.group(3)}") if m else None


def valor_a_hm(v):
    """Convierte un encabezado/celda de hora a (h, m) o None."""
    if isinstance(v, datetime):
        return v.hour, v.minute
    if isinstance(v, time):
        return v.hour, v.minute
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::\d{2})?\s*", v)
        if m:
            return int(m.group(1)), int(m.group(2))
    return None


def horas_a_timestamps(fecha, horas):
    """
    Lista de valores de hora (en orden de fila/columna) -> timestamps.
    Maneja '00:00' como ultimo bloque (dia siguiente) y '23:59' (redondea a 00:00 del dia siguiente).
    """
    out, dia = [], pd.Timestamp(fecha).normalize()
    prev = None
    for i, v in enumerate(horas):
        hm = valor_a_hm(v)
        if hm is None:
            out.append(pd.NaT)
            continue
        h, m = hm
        if m == 59:  # '23:59' -> 00:00 (se roll-ea abajo)
            h, m = (h + 1) % 24, 0
        if prev is not None and (h, m) <= prev:  # cruzo medianoche (p.ej. 23:30 -> 00:00)
            dia = dia + timedelta(days=1)
        out.append(dia + timedelta(hours=h, minutes=m))
        prev = (h, m)
    return out


# --------------------------------------------------------------------------- #
# B1  Mantenimientos
# --------------------------------------------------------------------------- #
B1_CRUDO = None


def fusionar_eventos(out):
    """
    Une filas CONTIGUAS (o solapadas) del mismo evento: misma ubicacion, equipo, tipo de equipo, tipo de
    mantenimiento, PROG y descripcion. Sirve cuando el reporte parte un evento largo en una fila por dia.
    No fusiona eventos distintos aunque sean del mismo equipo (distinta descripcion o separados en el tiempo).
    """
    claves = ["ubicacion", "equipo", "tipo_equipo", "tipo_mantenimiento", "prog", "descripcion"]
    o = out.copy()
    o["_k"] = o[claves].fillna("").astype(str).agg("|".join, axis=1)
    o = o.sort_values(["_k", "fecha_inicio", "fecha_fin"]).reset_index(drop=True)
    fin_prev = o.groupby("_k")["fecha_fin"].transform(lambda x: x.cummax().shift())
    nuevo = fin_prev.isna() | (o["fecha_inicio"] > fin_prev + pd.Timedelta(minutes=1))
    o["_ev"] = nuevo.cumsum()
    agg = o.groupby("_ev").agg(
        fecha_inicio=("fecha_inicio", "min"), fecha_fin=("fecha_fin", "max"),
        n_filas_origen=("fecha_inicio", "size"), **{c: (c, "first") for c in claves})
    nota(f"B1: fusion de filas contiguas del mismo evento: {len(out)} filas -> {len(agg)} eventos "
         f"({len(out) - len(agg)} filas absorbidas). Maximo de filas fusionadas en un evento: {int(agg['n_filas_origen'].max())}.")
    cols = ["fecha_inicio", "fecha_fin", "ubicacion", "equipo", "tipo_equipo", "tipo_mantenimiento",
            "prog", "descripcion", "n_filas_origen"]
    return agg[cols].sort_values(["fecha_inicio", "ubicacion", "equipo"]).reset_index(drop=True)


def b1_mantenimientos(ruta, fusionar=True):
    """`ruta` puede ser un .xlsx o una CARPETA con varios .xlsx (el portal limita cada consulta a 3 meses)."""
    ruta = Path(ruta)
    archivos = sorted(ruta.rglob("*.xlsx")) if ruta.is_dir() else [ruta]
    archivos = [x for x in archivos if not x.name.startswith("~$")]
    if not archivos:
        raise SystemExit(f"B1: no encontre .xlsx en {ruta}")
    frames = []
    for arch in archivos:
        xl = pd.ExcelFile(arch)
        hojas = [h for h in xl.sheet_names if "EJECUT" in norm(h)] or xl.sheet_names[:1]
        n_arch = 0
        for h in hojas:
            raw = xl.parse(h, header=None, dtype=object)
            fila_h = next((i for i in range(min(30, len(raw)))
                           if {"INICIO", "FINAL"} <= {norm(c) for c in raw.iloc[i].dropna()}), None)
            if fila_h is None:
                nota(f"B1: {arch.name} hoja '{h}' sin encabezado INICIO/FINAL; se omite.")
                continue
            d = raw.iloc[fila_h + 1:].copy()
            d.columns = [norm(c) for c in raw.iloc[fila_h]]
            d = d.loc[:, ~pd.Index(d.columns).isin(["NAN", ""])]
            n_arch += len(d.dropna(how="all"))
            frames.append(d)
        nota(f"B1: {arch.name}: {n_arch} filas crudas")
    df = pd.concat(frames, ignore_index=True).dropna(how="all")

    n0 = len(df)
    df = df.drop_duplicates()
    nota(f"B1: {n0} filas crudas -> {len(df)} tras deduplicar (factor ~{n0 / max(len(df), 1):.1f}x).")

    nota("B1: columnas encontradas en el Excel -> " + ", ".join(map(str, df.columns)))

    def col(*cands, obligatoria=True):
        for c in cands:
            if c in df.columns:
                return df[c]
        if obligatoria:
            raise KeyError(f"B1: no encuentro ninguna de {cands}. Columnas: {list(df.columns)}")
        nota(f"B1 AVISO: no existe ninguna columna {cands}; se deja vacia.")
        return pd.Series([None] * len(df), index=df.index)

    out = pd.DataFrame({
        "fecha_inicio": pd.to_datetime(col("INICIO"), dayfirst=True, errors="coerce"),
        "fecha_fin": pd.to_datetime(col("FINAL"), dayfirst=True, errors="coerce"),
        "ubicacion": col("UBICACION").astype(str).str.strip(),
        "equipo": col("EQUIPO").astype(str).str.strip(),
        "tipo_equipo": col("TIPO EQUIPO").astype(str).str.strip(),
        "tipo_mantenimiento": col("TIPO MANTENIMIENTO", "TIPO", obligatoria=False),
        "prog": col("PROG", obligatoria=False),
        "descripcion": col("DESCRIPCION").astype(str).str.strip(),
    })
    out["prog"] = out["prog"].astype(str).str.strip().str.upper().replace({"NONE": None, "NAN": None})
    malas = out["fecha_inicio"].isna() | out["fecha_fin"].isna()
    if malas.any():
        nota(f"B1: {int(malas.sum())} filas con fecha ilegible descartadas.")
        out = out[~malas]
    # traslape con la ventana
    out = out[(out["fecha_fin"] >= VENTANA_INI) & (out["fecha_inicio"] <= VENTANA_FIN)].copy()
    global B1_CRUDO
    dia24 = ((out["fecha_inicio"] == out["fecha_inicio"].dt.normalize())
             & (out["fecha_fin"] - out["fecha_inicio"] == pd.Timedelta(days=1)))
    nota(f"B1: {len(out)} filas en la ventana; {int(dia24.sum())} ({dia24.mean():.0%}) duran exactamente un dia "
         "calendario (00:00 a 00:00). Si el porcentaje es alto, el reporte parte los eventos largos en una fila por dia.")
    B1_CRUDO = out.copy()
    # cobertura: dias de la ventana sin NINGUNA fila (hueco entre consultas o dia sin datos)
    ini_d = out["fecha_inicio"].dt.normalize()
    fin_d = (out["fecha_fin"] - pd.Timedelta(minutes=1)).dt.normalize()
    cubiertos = set()
    for a_, b_ in set(zip(ini_d, fin_d)):
        cubiertos.update(pd.date_range(a_, max(a_, b_), freq="D"))
    todos = pd.date_range(VENTANA_INI.normalize(), VENTANA_FIN.normalize(), freq="D")
    faltan = [d for d in todos if d not in cubiertos]
    if faltan:
        nota(f"B1 COBERTURA: {len(faltan)} de {len(todos)} dias de la ventana no tienen ninguna fila. "
             f"Primeros: {', '.join(d.strftime('%Y-%m-%d') for d in faltan[:8])}"
             f"{' ...' if len(faltan) > 8 else ''}. Revisar si falta un tramo descargado.")
    else:
        nota(f"B1 COBERTURA: los {len(todos)} dias de la ventana tienen al menos una fila.")
    if fusionar:
        out = fusionar_eventos(out)
    out["duracion_horas"] = ((out["fecha_fin"] - out["fecha_inicio"]).dt.total_seconds() / 3600).round(2)
    out["instalacion_id"] = out["ubicacion"] + " | " + out["equipo"]
    nota("B1: conteo por PROG -> " + out["prog"].value_counts(dropna=False).to_dict().__repr__())
    for t in (out, B1_CRUDO):
        for c in ("fecha_inicio", "fecha_fin"):
            t[c] = t[c].dt.strftime(FMT)
    return out.reset_index(drop=True)


# --------------------------------------------------------------------------- #
# B2  Potencia media horaria indisponible (UG / LLTT / SSEE)
# --------------------------------------------------------------------------- #
def tipo_hoja(nombre):
    n = norm(nombre)
    for t in ("LLTT", "SSEE", "UG"):
        if t in n:
            return t
    return None


def b2_potencia(carpeta, solo_no_cero=False, vacio_a_cero=False):
    """
    Formato real (validado con captura del 25/11/2024): por hoja UG/LLTT/SSEE,
    encabezado con FECHA, SEMANA OPERATIVA, EMPRESA, UBICACION, EQUIPO, TIPO y luego 48 bloques (0:30 ... 0:00).
    La fecha viene en cada fila. Un equipo puede repetirse con distinto TIPO (CORR./PREV./PRUEBA/OTROS).
    """
    frames = []
    archivos = sorted(Path(carpeta).rglob("*.xlsx"))
    if not archivos:
        nota(f"B2: sin .xlsx en {carpeta}")
    base = pd.Timestamp("2000-01-01")
    for f in archivos:
        fecha_arch = parse_fecha_archivo(f.name)
        wb = load_workbook(f, read_only=True, data_only=True)
        for ws in wb.worksheets:
            tipo = tipo_hoja(ws.title)
            if tipo is None:
                continue
            filas = list(ws.iter_rows(values_only=True))
            ih = next((i for i, r in enumerate(filas[:40])
                       if sum(valor_a_hm(c) is not None for c in r) >= 40), None)
            if ih is None:
                nota(f"B2: {f.name}/{ws.title}: no hallo encabezado de 48 bloques.")
                continue
            enc = filas[ih]
            idx_h = [j for j, c in enumerate(enc) if valor_a_hm(c) is not None]
            nombres = {norm(c): j for j, c in enumerate(enc) if j not in idx_h and c is not None}
            offs = [t - base for t in horas_a_timestamps(base, [enc[j] for j in idx_h])]
            j_fecha = nombres.get("FECHA")
            j_id = [nombres[k] for k in ("UBICACION", "EQUIPO") if k in nombres]
            if not j_id:
                resto = [j for k, j in nombres.items()
                         if k not in ("FECHA", "SEMANA OPERATIVA", "EMPRESA", "TIPO")]
                j_id = resto[:1]
            if not j_id or (j_fecha is None and fecha_arch is None):
                nota(f"B2: {f.name}/{ws.title}: faltan columnas de instalacion o fecha; se omite.")
                continue
            j_tipo, j_emp = nombres.get("TIPO"), nombres.get("EMPRESA")
            reg = []
            for r in filas[ih + 1:]:
                get = lambda j: (r[j] if j is not None and j < len(r) else None)
                inst = " | ".join(str(get(j)).strip() for j in j_id if get(j) not in (None, ""))
                if not inst:
                    continue
                fch = pd.to_datetime(get(j_fecha), dayfirst=True, errors="coerce") if j_fecha is not None else fecha_arch
                if pd.isna(fch):
                    continue
                reg.append([fch.normalize(), inst, str(get(j_emp) or "").strip(),
                            str(get(j_tipo) or "").strip()] + [get(j) for j in idx_h])
            if not reg:
                continue
            wide = pd.DataFrame(reg, columns=["fecha", "instalacion", "empresa", "tipo_mantenimiento"] + list(range(len(idx_h))))
            largo = wide.melt(id_vars=["fecha", "instalacion", "empresa", "tipo_mantenimiento"],
                              var_name="bloque", value_name="potencia_indisponible_mw")
            largo["fecha_hora"] = largo["fecha"] + largo["bloque"].map(dict(enumerate(offs)))
            largo["tipo_instalacion"] = tipo
            largo["potencia_indisponible_mw"] = pd.to_numeric(largo["potencia_indisponible_mw"], errors="coerce")
            frames.append(largo.drop(columns=["fecha", "bloque"]))
    cols = ["fecha_hora", "instalacion", "tipo_instalacion", "tipo_mantenimiento", "empresa",
            "potencia_indisponible_mw"]
    if not frames:
        return pd.DataFrame(columns=cols)
    df = pd.concat(frames, ignore_index=True)
    df = df[df["fecha_hora"].notna() & en_ventana(df["fecha_hora"])]
    resumen = df.groupby("tipo_instalacion")["potencia_indisponible_mw"].agg(
        filas="size", vacias=lambda s: int(s.isna().sum()),
        no_cero=lambda s: int((s.fillna(0) != 0).sum()), maximo_mw="max")
    nota("B2: resumen por tipo de instalacion:\n" + resumen.to_string())
    if vacio_a_cero:
        df["potencia_indisponible_mw"] = df["potencia_indisponible_mw"].fillna(0)
        nota("B2: --vacio-a-cero activo: celdas vacias convertidas a 0 (supuesto: vacio = sin indisponibilidad).")
    else:
        nota("B2: las celdas VACIAS del Excel se dejan vacias (no se asumio que sean 0). "
             "Confirmar su significado antes de modelar; ver opcion --vacio-a-cero.")
    if solo_no_cero:
        df = df[df["potencia_indisponible_mw"].fillna(0) != 0]
        nota("B2: --solo-no-cero activo: se descartaron filas en cero/vacias.")
    df = df.drop_duplicates(["fecha_hora", "instalacion", "tipo_instalacion", "tipo_mantenimiento"])
    df["fecha_hora"] = df["fecha_hora"].dt.strftime(FMT)
    return df[cols].reset_index(drop=True)


# --------------------------------------------------------------------------- #
# B3  Reprograma Diario de Operacion
# --------------------------------------------------------------------------- #
def a_numero(txt):
    s = re.sub(r"[^\d,.\-]", "", str(txt))
    if not s:
        return None
    # formato peruano/ingles: coma = miles, punto = decimal
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def texto_tras_etiqueta(filas, clave, filas_siguientes=0):
    """
    Busca una celda que contenga `clave`. Devuelve, en orden de prioridad:
      1) el texto tras ':' en la misma celda,
      2) la siguiente celda no vacia de la misma fila,
      3) (si filas_siguientes > 0) la primera celda no vacia de las filas de abajo.
    """
    for i, r in enumerate(filas):
        for j, c in enumerate(r):
            if c is not None and clave in norm(c):
                tras = str(c).split(":", 1)[1].strip() if ":" in str(c) else ""
                if tras:
                    return tras
                for k in r[j + 1:]:
                    if k is not None and str(k).strip():
                        return k
                for r2 in filas[i + 1:i + 1 + filas_siguientes]:
                    for k in r2:
                        if k is not None and str(k).strip():
                            return k
                return None
    return None


def limpiar_motivo(x):
    """Texto en una linea, sin caracteres iniciales (- = + @ viñetas) que Excel interpretaria como formula."""
    if x is None:
        return None
    t = " ".join(str(x).split())
    t = re.sub(r"^[\s\-=+@\u2022\u2013\u2014*]+", "", t)
    return t or None


def categoria_motivo(motivo):
    """
    Clasifica el motivo en texto libre. Puede devolver varias categorias unidas con '+'
    (ej. 'hidrico+demanda'). Es una lectura por palabras clave: revisar a mano los casos 'otro'.
    """
    if motivo is None or (isinstance(motivo, float) and pd.isna(motivo)):
        return None
    t = norm(motivo)
    cats = []
    if any(k in t for k in ("INDISPONIBILIDAD", "FALLA", "AVERIA", "DISPARO", "RESTRICCION", "MANTENIMIENTO")):
        cats.append("falla_indisponibilidad")
    if "EOLIC" in t or "VIENTO" in t:
        cats.append("eolico")
    if "HIDRIC" in t or "HIDRAULIC" in t or "CAUDAL" in t:
        cats.append("hidrico")
    if "DEMANDA" in t:
        cats.append("demanda")
    return "+".join(cats) if cats else "otro"


def b3_reprograma(carpeta):
    filas_out = []
    for f in sorted(Path(carpeta).rglob("*.xlsx")):
        if f.name.upper().startswith("RS_"):
            continue  # reserva secundaria: prioridad menor, fuera del alcance de B3
        if not re.match(r"(?i)reprog", f.name):
            continue
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = list(ws.iter_rows(min_row=1, max_row=250, values_only=True))
        # letra: 'REPROGRAMA A' en la cabecera; si no, del nombre del archivo / carpeta
        letra = None
        for r in filas[:8]:
            for c in r:
                m = re.fullmatch(r"\s*REPROGRAMA\s+([A-Za-z])\s*", str(c or ""), flags=re.I)
                if m:
                    letra = m.group(1).upper()
        if letra is None:
            m = re.search(r"([A-Za-z])\.xlsx$", f.name)
            letra = m.group(1).upper() if m else (f.parent.name.upper() if len(f.parent.name) == 1 else None)
        # fecha: nombre de hoja, o celda dd/mm/aaaa de la cabecera, o nombre de archivo
        fecha = next((parse_fecha_archivo(w.title) for w in wb.worksheets if parse_fecha_archivo(w.title) is not None), None)
        if fecha is None:
            for r in filas[:8]:
                for c in r:
                    m = re.fullmatch(r"\s*(\d{2})/(\d{2})/(20\d{2})\s*", str(c or ""))
                    if m:
                        fecha = pd.Timestamp(f"{m.group(3)}-{m.group(2)}-{m.group(1)}")
        fecha = fecha if fecha is not None else parse_fecha_archivo(f.name)
        # tramo horario cubierto: filas bajo la cabecera 'Dia Hora' en la columna de horas
        i_cab = next((i for i, r in enumerate(filas[:15]) if any(norm(c) == "DIA HORA" for c in r if c)), None)
        ini = fin = None
        n_bloques = 0
        if i_cab is not None and fecha is not None:
            j_h = next(j for j, c in enumerate(filas[i_cab]) if c and norm(c) == "DIA HORA")
            horas = []
            for r in filas[i_cab + 1:]:
                v = r[j_h] if j_h < len(r) else None
                if valor_a_hm(v) is None:
                    break
                horas.append(v)
            if horas:
                tss = horas_a_timestamps(fecha, horas)
                ini, fin, n_bloques = tss[0], tss[-1], len(tss)
        motivo = texto_tras_etiqueta(filas, "MOTIVO", filas_siguientes=3)
        costo = texto_tras_etiqueta(filas, "COSTO TOTAL")
        if fecha is None or letra is None:
            nota(f"B3: {f.name}: fecha={fecha}, letra={letra}; revisar a mano.")
        motivo_limpio = limpiar_motivo(motivo)
        filas_out.append({
            "fecha": fecha.strftime("%Y-%m-%d") if fecha is not None else None,
            "letra_reprograma": letra,
            "motivo": motivo_limpio,
            "motivo_categoria": categoria_motivo(motivo_limpio),
            "costo_total_soles": a_numero(costo) if costo is not None else None,
            "fecha_hora_inicio": ini.strftime(FMT) if ini is not None else None,
            "fecha_hora_fin": fin.strftime(FMT) if fin is not None else None,
            "n_bloques": n_bloques,
            "horas_cubiertas": n_bloques * 0.5,
            "archivo": f.name})
    cols = ["fecha", "letra_reprograma", "motivo", "motivo_categoria", "costo_total_soles",
            "fecha_hora_inicio", "fecha_hora_fin", "n_bloques", "horas_cubiertas", "archivo"]
    df = pd.DataFrame(filas_out, columns=cols)
    if df.empty:
        return df
    df = df.drop_duplicates(["fecha", "letra_reprograma"])
    df = df[pd.to_datetime(df["fecha"]).between(VENTANA_INI.normalize(), VENTANA_FIN)]
    for _, r in df.iterrows():
        nota(f"B3: {r['fecha']} {r['letra_reprograma']} | {r['fecha_hora_inicio']} -> {r['fecha_hora_fin']} | "
             f"S/ {r['costo_total_soles']} | categoria: {r['motivo_categoria']} | motivo: {r['motivo']}")
    sin_m, sin_c = int(df["motivo"].isna().sum()), int(df["costo_total_soles"].isna().sum())
    if sin_m or sin_c:
        nota(f"B3: sin motivo en {sin_m} reprogramas y sin costo en {sin_c} (revisar esas celdas a mano).")
    nota("B3 (recordatorio): costo_total_soles es variable de IMPACTO/negocio (ahorro potencial), "
         "NO es el label del modelo ni una feature de entrenamiento. Probablemente es el costo del tramo "
         "restante del dia (baja con horas_cubiertas): no comparar letras entre si sin normalizar por horas.")
    nota("B3: motivo_categoria es una clasificacion por palabras clave (falla_indisponibilidad / eolico / "
         "hidrico / demanda; puede combinar varias con '+'). Revisar a mano los 'otro'.")
    return df.sort_values(["fecha", "letra_reprograma"]).reset_index(drop=True)


# --------------------------------------------------------------------------- #
# B4  PR06 (flag RDO = label)
# --------------------------------------------------------------------------- #
def fecha_y_letra_pr06(ruta):
    """
    Fecha y letra de un archivo PR06. Letra = carpeta de 1 letra (A,B,C...) o sufijo del nombre
    (indices2511B.xlsx / indices20241125B.xlsx). Fecha = AAAAMMDD en el nombre; si no, ddmm + anio de la ruta;
    si no, carpetas anio/mes/dia numericas.
    """
    ruta = Path(ruta)
    letra = None
    if len(ruta.parent.name) == 1 and ruta.parent.name.isalpha():
        letra = ruta.parent.name.upper()
    m = re.search(r"(?i)\d{4,8}[_\-]?([A-Z])\.xlsx$", ruta.name)
    if m and letra is None:
        letra = m.group(1).upper()
    fecha = parse_fecha_archivo(ruta.name)
    anio = next((p for p in ruta.parts if re.fullmatch(r"20\d{2}", p)), None)
    if fecha is None and anio:
        m = re.search(r"(\d{2})(\d{2})[A-Za-z]?\.xlsx$", ruta.name)
        if m:
            try:
                fecha = pd.Timestamp(f"{anio}-{m.group(2)}-{m.group(1)}")
            except ValueError:
                fecha = None
    if fecha is None and anio:
        nums = [int(p) for p in ruta.parts[ruta.parts.index(anio) + 1:-1] if re.fullmatch(r"\d{1,2}", p)]
        if len(nums) >= 2:
            try:
                fecha = pd.Timestamp(year=int(anio), month=nums[0], day=nums[1])
            except ValueError:
                fecha = None
    return fecha, letra


def diagnostico_abg(por_archivo, rep):
    """
    Compara alfa/beta/gamma entre letras consecutivas del mismo dia y los confronta con la hora de
    inicio de cada reprograma (B3).
      * cambian_antes_del_inicio == 0  -> solo cambian bloques futuros: pronosticos recalculados.
      * cambian_antes_del_inicio  > 0  -> tambien cambian bloques previos al inicio: revisar con cuidado.
    """
    pa = por_archivo.copy()
    pa["fecha_hora"] = pd.to_datetime(pa["fecha_hora"])
    inicio = {}
    if rep is not None and not rep.empty:
        for r in rep.itertuples():
            if pd.notna(r.fecha_hora_inicio):
                inicio[(r.fecha, r.letra_reprograma)] = pd.Timestamp(r.fecha_hora_inicio)
    filas = []
    for dia, g in pa.groupby("id_dia"):
        fecha = f"{dia[:4]}-{dia[4:6]}-{dia[6:]}"
        letras = sorted(g["letra_origen"].dropna().unique())
        for ant, nue in zip(letras, letras[1:]):
            a = g[g["letra_origen"] == ant].drop_duplicates("fecha_hora").set_index("fecha_hora")
            b = g[g["letra_origen"] == nue].drop_duplicates("fecha_hora").set_index("fecha_hora")
            comunes = a.index.intersection(b.index)
            ini = inicio.get((fecha, nue))
            for c in ("alfa", "beta", "gamma"):
                if c not in pa.columns:
                    continue
                x, y = a.loc[comunes, c], b.loc[comunes, c]
                dif = comunes[~((x == y) | (x.isna() & y.isna()))]
                filas.append({
                    "dia": dia, "de_letra": ant, "a_letra": nue, "variable": c,
                    "n_bloques_comunes": len(comunes),
                    "n_cambian": len(dif),
                    "primer_cambio": dif.min().strftime(FMT) if len(dif) else None,
                    "ultimo_cambio": dif.max().strftime(FMT) if len(dif) else None,
                    "inicio_reprog_b3": ini.strftime(FMT) if ini is not None else None,
                    "cambian_antes_del_inicio": int((dif < ini).sum()) if ini is not None else None,
                })
    return pd.DataFrame(filas)


def hora_a_timestamps_pr06(valores, fecha):
    """HORA del PR06: normalmente una fecha-hora completa (datetime o texto); a veces solo 'HH:MM'."""
    serie = pd.Series(list(valores), dtype=object)
    try:
        directo = pd.to_datetime(serie, errors="coerce", format="mixed")
    except (TypeError, ValueError):
        directo = pd.to_datetime(serie, errors="coerce")
    ok = directo.notna() & ((directo.dt.normalize() - pd.Timestamp(fecha).normalize()).abs() <= pd.Timedelta(days=1))
    respaldo = horas_a_timestamps(fecha, list(valores))
    return [d if o else r for d, o, r in zip(directo, ok, respaldo)]


def b4_pr06(carpeta, mant, rep=None):
    frames = []
    for f in sorted(Path(carpeta).rglob("*.xlsx")):
        fecha, letra = fecha_y_letra_pr06(f)
        if fecha is None:
            nota(f"B4: no pude sacar la fecha de '{f}'; se omite.")
            continue
        try:
            raw = pd.read_excel(f, header=None, dtype=object)
            ih = next((i for i in range(min(30, len(raw)))
                       if "RDO" in {norm(c) for c in raw.iloc[i].dropna()}), None)
            if ih is None:
                nota(f"B4: {f.name}: no hallo columna RDO.")
                continue
            df = raw.iloc[ih + 1:].copy()
            df.columns = [norm(c).lower() for c in raw.iloc[ih]]
            # columnas sin nombre ("nan"/vacias) o repetidas rompen el concat de miles de archivos: se quitan
            df = df.loc[:, [c not in ("nan", "") for c in df.columns]]
            df = df.loc[:, ~df.columns.duplicated()]
            df = df.dropna(subset=["rdo"])
            col_hora = next((c for c in df.columns if c.startswith("hora")), df.columns[0])
            df["fecha_hora"] = hora_a_timestamps_pr06(df[col_hora].tolist(), fecha)
            malas = df["fecha_hora"].isna()
            if malas.any():
                nota(f"B4 AVISO: {f.name}: {int(malas.sum())} fila(s) con hora ilegible descartadas "
                     f"(ej.: {list(df.loc[malas, col_hora].astype(str).unique()[:3])}).")
                df = df[~malas].copy()
            if df.empty:
                nota(f"B4 AVISO: {f.name}: sin filas validas; se omite.")
                continue
            if len(df) != 48:
                nota(f"B4 AVISO: {f.name}: tiene {len(df)} filas validas en lugar de 48.")
            df["rdo"] = pd.to_numeric(df["rdo"], errors="coerce").fillna(0).astype(int)
            df["id_dia"] = fecha.strftime("%Y%m%d")
            df["letra_origen"] = letra
            df["archivo"] = str(f.relative_to(carpeta))
            # --- id_reprograma: racha contigua de RDO=1 (fallback; se mejora abajo con la primera letra) ---
            run = (df["rdo"] != df["rdo"].shift()).cumsum()
            orden = {r: n + 1 for n, r in enumerate(run[df["rdo"] == 1].drop_duplicates())}
            if letra:
                df["id_reprograma"] = [f"{d}{letra}" if v == 1 else None for d, v in zip(df["id_dia"], df["rdo"])]
                df["id_fuente"] = ["carpeta_archivo" if v == 1 else None for v in df["rdo"]]
            else:
                df["id_reprograma"] = [f"{d}_R{orden[r]}" if v == 1 else None
                                       for d, r, v in zip(df["id_dia"], run, df["rdo"])]
                df["id_fuente"] = ["racha_rdo" if v == 1 else None for v in df["rdo"]]
            tramos = []
            for _, g in df[df["rdo"] == 1].groupby(run[df["rdo"] == 1]):
                tramos.append(f"{g['fecha_hora'].iloc[0]:%H:%M}-{g['fecha_hora'].iloc[-1]:%H:%M}")
            nota(f"B4: {f.name} (letra={letra}): {len(df)} filas, {int((df['rdo'] == 1).sum())} bloques RDO=1, "
                 f"tramos: {', '.join(tramos) if tramos else 'ninguno'}")
            frames.append(df)
        except Exception as e:  # un archivo raro no debe tumbar las 30 minutos de proceso
            nota(f"B4 ERROR en {f}: {type(e).__name__}: {e}. Se omite ese archivo.")
    if not frames:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()
    frames = [fr.loc[:, ~fr.columns.duplicated()].reset_index(drop=True) for fr in frames]
    df = pd.concat(frames, ignore_index=True)
    df = df[df["fecha_hora"].notna() & en_ventana(df["fecha_hora"])].copy()
    for c in ("alfa", "beta", "gamma"):
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    df["letra_orden"] = df["letra_origen"].fillna("~")
    df = df.sort_values(["fecha_hora", "letra_orden"]).reset_index(drop=True)
    por_archivo = df.drop(columns=["letra_orden"]).copy()

    # --- la letra A no siempre tiene archivo de indices (en 25/11/2024 solo hay B, C, D) ---
    sin_a = [dia for dia, gd in df.groupby("id_dia")
             if gd["letra_origen"].dropna().nunique() and "A" not in set(gd["letra_origen"].dropna())]
    if sin_a:
        nota(f"B4: {len(sin_a)} dia(s) sin archivo de indices de la letra A ({', '.join(sin_a[:5])}"
             f"{'...' if len(sin_a) > 5 else ''}). Se confirmo que COES no lo publica para 20241125, asi que NO es un dato "
             "faltante: la primera letra disponible es la primera version existente. Los bloques RDO=1 que ya aparecen "
             "en ella se asignan a esa letra por regla (id_fuente='primera_letra_donde_aparece'); si hay un bloque "
             "aislado al inicio del dia (p.ej. 00:30) su atribucion es ambigua. Para el split usar id_dia.")

    # --- los archivos del mismo dia son ACUMULATIVOS: consolidar a una fila por fecha_hora ---
    g = df.groupby("fecha_hora")
    n_multi = int((g.size() > 1).sum())
    if n_multi:
        ultimo = g["rdo"].last()
        perdidos = int(((g["rdo"].max() == 1) & (ultimo == 0)).sum())
        nota(f"B4: {n_multi} bloques aparecen en varios archivos del mismo dia. Se consolida a UNA fila por fecha_hora "
             "usando el archivo de la letra mas alta (el reprograma mas reciente). Los archivos parecen ACUMULATIVOS "
             "(cada letra agrega los bloques RDO=1 que la preceden). "
             f"Bloques con RDO=1 en un archivo anterior pero RDO=0 en el mas reciente (romperian lo acumulativo): {perdidos}.")
        for c in ("alfa", "beta", "gamma"):
            if c in df.columns:
                dif = g[c].nunique(dropna=False) > 1
                if dif.any():
                    pos = dif[dif].index
                    nota(f"B4: bloques donde {c} difiere entre archivos: {int(dif.sum())} "
                         f"(de {pos.min():%Y-%m-%d %H:%M} a {pos.max():%Y-%m-%d %H:%M}). "
                         "Ver B4_diagnostico_abg.csv: si empiezan despues de la hora de emision de cada "
                         "reprograma, son pronosticos recalculados.")
                else:
                    nota(f"B4: {c} es identico en todos los archivos.")

    # valores de alfa/beta/gamma del PRIMER archivo del dia (alternativa conservadora contra fuga de datos)
    abg = [c for c in ("alfa", "beta", "gamma") if c in df.columns]
    primero = df.drop_duplicates("fecha_hora", keep="first").set_index("fecha_hora")[abg]

    primera = (df[df["rdo"] == 1].dropna(subset=["letra_origen"])
               .sort_values("letra_orden").groupby("fecha_hora")["letra_origen"].first())
    n_arch = df[df["rdo"] == 1].groupby("fecha_hora").size()
    df = g.tail(1).copy().reset_index(drop=True)  # ya ordenado: la ultima fila = letra mas alta
    m = (df["rdo"] == 1) & df["fecha_hora"].isin(primera.index)
    df.loc[m, "id_reprograma"] = df.loc[m, "id_dia"] + df.loc[m, "fecha_hora"].map(primera)
    df.loc[m, "id_fuente"] = "primera_letra_donde_aparece"
    df["n_archivos_con_rdo1"] = df["fecha_hora"].map(n_arch).fillna(0).astype(int)
    for c in abg:
        df[f"{c}_primer_archivo"] = df["fecha_hora"].map(primero[c])

    nota("B4: id_reprograma = AAAAMMDD + la PRIMERA letra cuyo archivo contiene ese bloque con RDO=1 "
         "(id_fuente='primera_letra_donde_aparece'). Es consistente con que cada tramo RDO precede al reprograma "
         "que lo incorpora, pero es una lectura de los datos: confirmarla en el PR-06. "
         "Si el archivo no indica letra se usa la racha contigua (id_fuente='racha_rdo'). "
         "Para el split train/test agrupar por id_dia (mas seguro) o por id_reprograma.")
    nota("B4: alfa, beta, gamma y CHM (de la letra mas alta del dia) se entregan sin descartar, junto con "
         "alfa/beta/gamma_primer_archivo (valores del primer archivo del dia). OJO: los de la letra mas alta pueden "
         "contener informacion posterior al bloque (fuga de datos); definir con B4_diagnostico_abg.csv cual usar. "
         "Segun el texto del PR-06 (Osinergmin 012-2021-OS/CD) "
         "alfa = parametro de generacion hidraulica, beta = generacion termica, gamma = demanda, calculados cada 30 min. "
         "Falta confirmar en el Anexo 1 del PR-06 como se interpreta el valor y las tolerancias.")

    # --- consistencia contra B1 ---
    forz = mant[mant["prog"] == "FORZADO"].copy()
    fi = pd.to_datetime(forz["fecha_inicio"]).to_numpy()
    ff = pd.to_datetime(forz["fecha_fin"]).to_numpy()

    def hay_forzado(ts):  # FORZADO en curso durante el bloque [t-30min, t]
        t0, t1 = (ts - timedelta(minutes=30)).to_datetime64(), ts.to_datetime64()
        return bool(((fi < t1) & (ff > t0)).any())

    def forzado_inicia_3h(ts):  # FORZADO que EMPEZO en las 3 h previas al bloque
        t0, t1 = (ts - timedelta(hours=3)).to_datetime64(), ts.to_datetime64()
        return bool(((fi > t0) & (fi <= t1)).any())

    df["forzado_activo_b1"] = df["fecha_hora"].apply(hay_forzado)
    df["forzado_inicio_3h_previas_b1"] = df["fecha_hora"].apply(forzado_inicia_3h)
    uno, cero = df[df["rdo"] == 1], df[df["rdo"] == 0]

    def pct(x):
        return f"{x.mean():.0%}" if len(x) else "n/a"

    nota(f"B4 consistencia (criterio laxo): de {len(uno)} bloques RDO=1, {int(uno['forzado_activo_b1'].sum())} tienen "
         f"un FORZADO en curso; entre los {len(cero)} bloques RDO=0 lo tiene {pct(cero['forzado_activo_b1'])}. "
         "Si ambos porcentajes son parecidos, este criterio no distingue nada (eventos FORZADO de todo el dia).")
    nota(f"B4 consistencia (criterio fino): de {len(uno)} bloques RDO=1, {int(uno['forzado_inicio_3h_previas_b1'].sum())} "
         f"({pct(uno['forzado_inicio_3h_previas_b1'])}) tienen un FORZADO que EMPEZO en las 3 h previas; "
         f"entre RDO=0: {pct(cero['forzado_inicio_3h_previas_b1'])}. No se forzo ninguna coincidencia.")
    dias = uno.groupby("id_dia").agg(
        bloques_rdo1=("rdo", "size"),
        con_forzado_en_curso=("forzado_activo_b1", "sum"),
        con_forzado_iniciado_3h=("forzado_inicio_3h_previas_b1", "sum"))
    dias["sin_forzado_iniciado_3h"] = dias["bloques_rdo1"] - dias["con_forzado_iniciado_3h"]
    dias["pct_forzado_en_curso_rdo0"] = dias.index.map(cero.groupby("id_dia")["forzado_activo_b1"].mean().round(3))
    dias["pct_forzado_iniciado_3h_rdo0"] = dias.index.map(
        cero.groupby("id_dia")["forzado_inicio_3h_previas_b1"].mean().round(3))

    for t in (df, por_archivo):
        t["fecha_hora"] = pd.to_datetime(t["fecha_hora"]).dt.strftime(FMT)
    extra = [c for c in ("alfa", "beta", "gamma", "chm") if c in df.columns]
    extra_primero = [f"{c}_primer_archivo" for c in abg]
    keep = ["fecha_hora", "id_dia", "id_reprograma", "id_fuente", "n_archivos_con_rdo1", "rdo"] + extra + \
           extra_primero + ["forzado_activo_b1", "forzado_inicio_3h_previas_b1", "archivo"]
    out = df[keep].rename(columns={"rdo": "label_rdo"})
    keep_pa = ["fecha_hora", "id_dia", "letra_origen", "archivo", "rdo"] + extra
    return out, dias.reset_index(), por_archivo[keep_pa].rename(columns={"rdo": "label_rdo"})


# --------------------------------------------------------------------------- #
# Tabla maestra de B (una fila por bloque de 30 min)
# --------------------------------------------------------------------------- #
def b_maestro(b4, b1, b2=None, rep=None):
    """
    Une el label (B4) con features sin fuga de datos (B1, B2) y columnas de IMPACTO (B3).
      * features seguras: n_forzados_activos_b1, forzado_*_b1, mw_indisp_ug/lltt/ssee
      * columnas 'impacto_*' (B3): solo para el calculo de ahorro, NO entrenar con ellas
      * alfa/beta/gamma/chm y sus *_primer_archivo: entregados pero SIN confirmar (posible fuga)
    """
    m = b4.copy()
    ts = pd.to_datetime(m["fecha_hora"])

    forz = b1[b1["prog"] == "FORZADO"]
    fi = pd.to_datetime(forz["fecha_inicio"]).to_numpy()
    ff = pd.to_datetime(forz["fecha_fin"]).to_numpy()

    def n_activos(t):
        t1, t0 = t.to_datetime64(), (t - timedelta(minutes=30)).to_datetime64()
        return int(((fi < t1) & (ff > t0)).sum())

    m["n_forzados_activos_b1"] = ts.apply(n_activos)

    if b2 is not None and not b2.empty:
        p = (b2.groupby(["fecha_hora", "tipo_instalacion"])["potencia_indisponible_mw"]
               .sum(min_count=1).unstack())
        p.columns = [f"mw_indisp_{str(c).lower()}" for c in p.columns]
        m = m.merge(p.reset_index(), on="fecha_hora", how="left")
        cols_b2 = [c for c in m.columns if c.startswith("mw_indisp_")]
        sin_b2 = int(m[cols_b2].isna().all(axis=1).sum())
        nota(f"MAESTRO: {sin_b2} de {len(m)} bloques no tienen dato de B2 (dia sin archivo de potencia o celda vacia); "
             "quedan vacios, no se asumio 0.")
    else:
        nota("MAESTRO: sin B2; la tabla maestra no incluye mw_indisp_*.")

    if rep is not None and not rep.empty:
        r = rep.sort_values(["fecha", "letra_reprograma"]).copy()
        r["id_dia"] = r["fecha"].str.replace("-", "", regex=False)
        imp = r.groupby("id_dia").agg(
            impacto_n_reprogramas_dia=("letra_reprograma", "nunique"),
            impacto_costo_soles_primera_letra=("costo_total_soles", "first"),
            impacto_costo_soles_ultima_letra=("costo_total_soles", "last")).reset_index()
        nota("MAESTRO: impacto_costo_soles_primera_letra cubre el dia completo (si la letra A empieza 00:30); "
             "impacto_costo_soles_ultima_letra solo cubre el tramo final del dia, no es el costo del dia.")
        m = m.merge(imp, on="id_dia", how="left")
    else:
        nota("MAESTRO: sin B3; la tabla maestra no incluye columnas impacto_*.")

    nota("MAESTRO: columnas impacto_* (B3) son para estimar ahorro, NO son features. "
         "alfa/beta/gamma/chm (y *_primer_archivo) se entregan pero quedan FUERA de la v1 del modelo "
         "hasta confirmar su significado y hora de emision. Split train/test por id_dia.")
    return m


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--mant", required=True, help="RptMantenimiento.xlsx (B1), una CARPETA con varios .xlsx de 3 meses, o B1_mantenimientos.csv ya generado")
    ap.add_argument("--pot-dir", help="carpeta con xlsx de potencia indisponible (B2)")
    ap.add_argument("--reprog-dir", help="carpeta raiz del reprograma (B3)")
    ap.add_argument("--pr06-dir", help="carpeta con indicesAAAAMMDD.xlsx (B4)")
    ap.add_argument("--out", default="salida_B")
    ap.add_argument("--sin-fusion", action="store_true", help="B1: NO fusionar filas contiguas del mismo evento")
    ap.add_argument("--solo-no-cero", action="store_true", help="B2: omitir filas en cero o vacias")
    ap.add_argument("--vacio-a-cero", action="store_true", help="B2: tratar celdas vacias como 0")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)

    if a.mant.lower().endswith(".csv"):
        b1 = pd.read_csv(a.mant)
        nota(f"B1: leido desde {a.mant} ({len(b1)} eventos ya fusionados); no se releen los Excel de mantenimientos.")
    else:
        b1 = b1_mantenimientos(a.mant, fusionar=not a.sin_fusion)
        b1.to_csv(out / "B1_mantenimientos.csv", index=False, encoding="utf-8-sig")
        B1_CRUDO.to_csv(out / "B1_mantenimientos_crudo.csv", index=False, encoding="utf-8-sig")
    b2 = None
    if a.pot_dir:
        b2 = b2_potencia(a.pot_dir, a.solo_no_cero, a.vacio_a_cero)
        b2.to_csv(out / "B2_potencia_indisponible.csv", index=False, encoding="utf-8-sig")
    rep = None
    if a.reprog_dir:
        rep = b3_reprograma(a.reprog_dir)
        rep.to_csv(out / "B3_reprograma_diario.csv", index=False, encoding="utf-8-sig")
    if a.pr06_dir:
        b4, cons, b4_pa = b4_pr06(a.pr06_dir, b1, rep)
        b4.to_csv(out / "B4_pr06_rdo.csv", index=False, encoding="utf-8-sig")
        b4_pa.to_csv(out / "B4_pr06_rdo_por_archivo.csv", index=False, encoding="utf-8-sig")
        cons.to_csv(out / "B4_consistencia_rdo_vs_forzado.csv", index=False, encoding="utf-8-sig")
        if not b4_pa.empty:
            try:
                diag = diagnostico_abg(b4_pa, rep)
                diag.to_csv(out / "B4_diagnostico_abg.csv", index=False, encoding="utf-8-sig")
                if not diag.empty and diag["cambian_antes_del_inicio"].notna().any():
                    antes = int(diag["cambian_antes_del_inicio"].fillna(0).sum())
                    nota(f"B4 diagnostico alfa/beta/gamma: {antes} cambios en bloques ANTERIORES al inicio del "
                         "reprograma que los recalcula (0 = solo cambian bloques futuros = pronosticos recalculados). "
                         "Detalle en B4_diagnostico_abg.csv.")
            except Exception as e:
                nota(f"B4 ERROR en el diagnostico alfa/beta/gamma ({type(e).__name__}: {e}); se omite, el resto sigue.")
    if a.pr06_dir and not b4.empty:
        try:
            b_maestro(b4, b1, b2, rep).to_csv(out / "B_dataset_maestro.csv", index=False, encoding="utf-8-sig")
        except Exception as e:
            nota(f"MAESTRO ERROR ({type(e).__name__}: {e}). B1, B3 y B4 ya estan guardados; avisar para corregir.")
    (out / "NOTAS_B.txt").write_text("\n\n".join(NOTAS), encoding="utf-8")
    print(f"\nListo. Archivos en {out.resolve()}")


if __name__ == "__main__":
    main()
